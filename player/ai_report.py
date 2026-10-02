"""The controller's own daily report with the AI's answers added next to it.

The controller uploads the report they filled in by hand («Отчет записи ДД.ММ.ГГГГ.xlsx»);
SyncVoice keeps the file as it is (colours, widths, comments) and adds columns
to the right: the AI's city / stations suggestion and its survey check, each
marked green where it agrees with the controller and red where it does not.
A second sheet sums it up. Only answers already saved are used — no Gemini requests.
"""
import io
import re
from datetime import datetime, timedelta

from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import AudioFile, CallAnalysis, CallComparison

MATCH_WINDOW = timedelta(minutes=5)
FIRST_AI_COLUMN = 17  # Q
HEADERS = {  # column titles searched in the header row (lower case, spaces collapsed)
    'phone': '№ телефона',
    'time': 'дата опроса',
    'interviewer': 'интервьюер',
    'city': 'город проживания',
    'listen_city': 'город слушания',
    'stations': 'радиостанции',
    'errors': 'ошибки',
    'comment': 'комментарий к ошибкам',
    'result': 'результат',
}
AI_COLUMNS = [
    ('ИИ: город проживания', 12.2),
    ('ИИ: город слушания', 10.5),
    ('ИИ: радиостанции / не слушал', 17.0),
    ('ИИ: ошибки', 8.5),
    ('ИИ: что нашёл при сверке анкеты', 45.0),
    ('Совпадение с контролёром', 18.0),
]

_thin = Side(style='thin')
_border = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
_header_fill = PatternFill(fill_type='solid', fgColor='FFDDEBF7')  # light blue: the AI part
_agree = PatternFill(fill_type='solid', fgColor='FFC6EFCE')  # light green
_disagree = PatternFill(fill_type='solid', fgColor='FFFFC7CE')  # light red
# The controllers' own colours: an error is orange, a recommendation is green.
_error = PatternFill(fill_type='solid', fgColor='FFFFC000')
_recommendation = PatternFill(fill_type='solid', fgColor='FF00FF00')

_NOISE = re.compile(r'\([^)]*\)|\bрадио\b|\bfm\b|\bфм\b')


def _norm_header(value) -> str:
    return ' '.join(str(value or '').casefold().split())


def _text(value) -> str:
    return str(value).strip() if value not in (None, '') else ''


def _phone_key(value) -> str:
    """Last 9 digits: the report keeps phones as numbers (375291234567.0)."""
    if isinstance(value, float):
        value = int(value)
    digits = re.sub(r'\D', '', str(value or ''))
    return digits[-9:]


def _city(value: str) -> str:
    return value.casefold().replace('ё', 'е').strip()


def stations_set(value: str) -> set:
    """Stations up to case, order, «Радио», «FM» and «(Другое)»; «слушала» = «слушал»."""
    out = set()
    for part in re.split(r'[,.;]', value or ''):
        part = _NOISE.sub(' ', part.casefold().replace('ё', 'е')).replace('слушала', 'слушал')
        part = ' '.join(part.split())
        if part:
            out.add(part)
    return out


def controller_error(errors: str, result: str) -> bool:
    return bool(errors.strip()) or result.casefold().startswith(('ошибк', 'брак'))


def find_header(ws):
    """(row number, {key: column number}) of the report's header row."""
    for row in ws.iter_rows(min_row=1, max_row=10):
        found = {}
        for cell in row:
            title = _norm_header(cell.value)
            for key, needle in HEADERS.items():
                if key not in found and title.startswith(needle.casefold()):
                    found[key] = cell.column
        if {'phone', 'time'} <= found.keys():
            return row[0].row, found
    raise ValueError('Не нашёл строку заголовков: нужны столбцы «№ телефона» и «Дата опроса Время».')


def _as_datetime(value):
    if isinstance(value, datetime):
        return value
    try:
        return datetime.fromisoformat(str(value).strip())
    except ValueError:
        return None


def _match(phone: str, started: datetime):
    """The SyncVoice call with this phone closest in time (within MATCH_WINDOW)."""
    if not phone or not started:
        return None
    started = timezone.make_aware(started) if timezone.is_naive(started) else started
    calls = (
        AudioFile.objects.filter(
            phone__endswith=phone,
            call_started_at__range=(started - MATCH_WINDOW, started + MATCH_WINDOW),
        ).select_related('analysis', 'comparison')
    )
    return min(calls, key=lambda a: abs(a.call_started_at - started), default=None)


def _discrepancy_lines(items):
    order = {'ошибка': 0, 'проверить': 1, 'рекомендация': 2}
    lines = []
    for d in sorted(items, key=lambda d: order.get(d.get('severity'), 3)):
        line = f"{d.get('severity', '')}: {d.get('field', '')} — анкета «{d.get('survey_value', '')}», " \
               f"разговор «{d.get('call_value', '')}»"
        if d.get('comment'):
            line += f". {d['comment']}"
        lines.append(line)
    return '\n'.join(lines)


class Totals:
    def __init__(self):
        self.rows = self.matched = 0
        self.missing = []
        self.city = [0, 0]  # agree, compared
        self.stations = [0, 0]
        self.errors = {'both': 0, 'ai_only': 0, 'missed': 0, 'clean': 0}
        self.models = set()


def build(file) -> tuple[bytes, Totals]:
    wb = load_workbook(file)
    ws = wb.worksheets[0]
    header_row, cols = find_header(ws)
    # The template has blank white-filled columns after the data; start after the last titled one,
    # but not before Q: in SyncVoice's own export the untitled P holds «исправлено».
    titled = [c.column for c in ws[header_row] if _text(c.value)]
    first = max(max(titled) + 1, FIRST_AI_COLUMN)

    for i, (title, width) in enumerate(AI_COLUMNS):
        cell = ws.cell(row=header_row, column=first + i, value=title)
        cell.font = Font(name='Calibri', size=11, bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = _border
        cell.fill = _header_fill
        ws.column_dimensions[get_column_letter(first + i)].width = width

    totals = Totals()
    for row in ws.iter_rows(min_row=header_row + 1):
        value = {key: row[col - 1].value for key, col in cols.items()}
        phone, started = _phone_key(value['phone']), _as_datetime(value['time'])
        if not phone or not started:
            continue
        city, stations = _text(value.get('city')), _text(value.get('stations'))
        errors, result = _text(value.get('errors')), _text(value.get('result'))
        checked = bool(city or stations or result)
        totals.rows += checked
        audio = _match(phone, started)
        cells = [ws.cell(row=row[0].row, column=first + i) for i in range(len(AI_COLUMNS))]
        for cell in cells:
            cell.font = Font(name='Calibri', size=11)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.border = _border
        if not audio:
            if checked:
                totals.missing.append(f"{started:%H:%M} {_text(value.get('interviewer'))}")
            cells[-1].value = 'нет в SyncVoice'
            continue
        totals.matched += checked
        notes = []

        suggestion = getattr(audio, 'analysis', None)
        if suggestion and suggestion.status == CallAnalysis.Status.DONE:
            totals.models.add(f'подсказка: {suggestion.model_name} / {suggestion.prompt_version}')
            cells[0].value = suggestion.city or None
            cells[1].value = suggestion.listen_city or None
            cells[2].value = suggestion.stations or None
            if checked:
                if city:
                    agree = _city(suggestion.city) == _city(city)
                    _mark(cells[0], agree, totals.city)
                    notes.append(f"город {'✓' if agree else '✗'}")
                if stations:
                    agree = stations_set(suggestion.stations) == stations_set(stations)
                    _mark(cells[2], agree, totals.stations)
                    notes.append(f"станции {'✓' if agree else '✗'}")

        check = getattr(audio, 'comparison', None)
        if check and check.status == CallComparison.Status.DONE:
            totals.models.add(f'сверка: {check.model_name} / {check.prompt_version}')
            items = check.discrepancies or []
            ai_error = any(d.get('severity') == 'ошибка' for d in items)
            cells[3].value = 'есть' if ai_error else None
            cells[4].value = _discrepancy_lines(items) or check.summary or None
            if ai_error:
                cells[4].fill = _error
            elif items:
                cells[4].fill = _recommendation
            if checked:
                mine = controller_error(errors, result)
                key = {(True, True): 'both', (True, False): 'ai_only',
                       (False, True): 'missed', (False, False): 'clean'}[ai_error, mine]
                totals.errors[key] += 1
                cells[3].fill = _agree if ai_error == mine else _disagree
                notes.append({'both': 'ошибка ✓', 'clean': 'без ошибок ✓',
                              'ai_only': 'ИИ: лишняя ошибка', 'missed': 'ИИ пропустил ошибку'}[key])
        elif check and check.status == CallComparison.Status.ERROR:
            cells[4].value = f'Сверка не удалась: {check.error}'
        cells[-1].value = ', '.join(notes) or ('не проверен контролёром' if not checked else None)

    if ws.auto_filter.ref:  # let the controller filter by the AI columns too
        ws.auto_filter.ref = f'A{header_row}:{get_column_letter(first + len(AI_COLUMNS) - 1)}{ws.max_row}'
    _summary_sheet(wb, totals)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue(), totals


def _mark(cell, agree: bool, counter: list):
    cell.fill = _agree if agree else _disagree
    counter[0] += agree
    counter[1] += 1


def _percent(part, whole) -> str:
    return f'{part} из {whole} ({round(100 * part / whole)} %)' if whole else '—'


def _missing(missing: list) -> str:
    if len(missing) > 15:
        return f"{len(missing)}: {', '.join(missing[:15])}… — скачайте этот день из АТС"
    return ', '.join(missing) or '—'


def summary_lines(totals: Totals) -> list[tuple[str, str]]:
    e = totals.errors
    flagged, real = e['both'] + e['ai_only'], e['both'] + e['missed']
    return [
        ('Проверено контролёром звонков', str(totals.rows)),
        ('Из них найдено в SyncVoice', str(totals.matched)),
        ('Не найдено в SyncVoice', _missing(totals.missing)),
        ('', ''),
        ('Город проживания: ИИ совпал с контролёром', _percent(*totals.city)),
        ('Радиостанции: ИИ совпал с контролёром', _percent(*totals.stations)),
        ('', ''),
        ('Ошибка и у контролёра, и у ИИ', str(e['both'])),
        ('Ошибка только у ИИ (лишняя)', str(e['ai_only'])),
        ('Ошибка только у контролёра (ИИ пропустил)', str(e['missed'])),
        ('Оба: без ошибок', str(e['clean'])),
        ('Ошибки ИИ подтверждены контролёром', _percent(e['both'], flagged)),
        ('Ошибки контролёра найдены ИИ', _percent(e['both'], real)),
        ('', ''),
        ('Модели и версии промптов', '\n'.join(sorted(totals.models)) or '—'),
    ]


def _summary_sheet(wb, totals: Totals):
    ws = wb.create_sheet('Итог ИИ')
    ws.column_dimensions['A'].width = 45
    ws.column_dimensions['B'].width = 60
    ws['A1'] = 'Сравнение ИИ с отчётом контролёра'
    ws['A1'].font = Font(name='Calibri', size=13, bold=True)
    for i, (title, value) in enumerate(summary_lines(totals), 3):
        ws.cell(row=i, column=1, value=title or None).font = Font(name='Calibri', size=11, bold=bool(title))
        cell = ws.cell(row=i, column=2, value=value or None)
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    note = ws.cell(row=ws.max_row + 2, column=1, value=(
        'В листе отчёта справа добавлены столбцы ИИ: зелёным — ИИ согласен с контролёром, '
        'красным — не согласен. В «что нашёл» оранжевым — ИИ считает ошибкой, '
        'ярко-зелёным — только замечания.'))
    note.alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=note.row, start_column=1, end_row=note.row, end_column=2)
    ws.row_dimensions[note.row].height = 48


def output_filename(name: str) -> str:
    stem = re.sub(r'\.xlsx$', '', name or 'Отчет', flags=re.I)
    return f'{stem} + ИИ.xlsx'
