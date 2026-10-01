import io
import http.client
import json
import shutil
import sys
import tempfile
import threading
import wave
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace
from http.server import ThreadingHTTPServer
from unittest import mock

import httpx
import openpyxl
import requests
from django import forms
from django.core.cache import cache
from django.core.files.base import ContentFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from google.genai import errors as genai_errors

from . import analysis, audio_paths, crm, crm_proxy, report, stations, worker
from . import compare as compare_module
from .analysis import SuggestedFields
from . import forms as forms_module
from .forms import SettingsForm
from .models import AppSettings, AudioFile, CallAnalysis, CallComparison, CallReview, RadioStation, Interviewer, PbxSync, Segment, parse_call_name
from . import pbx
from .pbx import FreePbxClient, PbxError, Recording, parse_cdr_html
from .sync import run_sync

MEDIA_ROOT = tempfile.mkdtemp()
DATA = bytes(range(256)) * 40  # 10240 bytes


@override_settings(MEDIA_ROOT=MEDIA_ROOT)
class StreamTests(TestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(MEDIA_ROOT, ignore_errors=True)

    def setUp(self):
        self.audio = AudioFile()
        self.audio.file.save('test.wav', ContentFile(DATA))
        self.url = reverse('player:stream', args=[self.audio.pk])

    def get(self, range_header=None):
        headers = {'Range': range_header} if range_header else {}
        return self.client.get(self.url, headers=headers)

    def test_full_file(self):
        r = self.get()
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r['Accept-Ranges'], 'bytes')
        self.assertEqual(b''.join(r.streaming_content), DATA)

    def test_range(self):
        r = self.get('bytes=100-199')
        self.assertEqual(r.status_code, 206)
        self.assertEqual(r['Content-Range'], f'bytes 100-199/{len(DATA)}')
        self.assertEqual(r.content, DATA[100:200])

    def test_open_range(self):
        r = self.get('bytes=10000-')
        self.assertEqual(r.status_code, 206)
        self.assertEqual(r.content, DATA[10000:])

    def test_suffix_range(self):
        r = self.get('bytes=-50')
        self.assertEqual(r.status_code, 206)
        self.assertEqual(r.content, DATA[-50:])

    def test_range_past_end_is_clamped(self):
        r = self.get('bytes=10200-99999')
        self.assertEqual(r['Content-Range'], f'bytes 10200-10239/{len(DATA)}')

    def test_unsatisfiable_range(self):
        r = self.get(f'bytes={len(DATA)}-')
        self.assertEqual(r.status_code, 416)
        self.assertEqual(r['Content-Range'], f'bytes */{len(DATA)}')

    def test_malformed_range_returns_full_file(self):
        r = self.get('items=0-10')
        self.assertEqual(r.status_code, 200)


@override_settings(MEDIA_ROOT=MEDIA_ROOT)
class SubtitlesTests(TestCase):
    def test_segments_in_order(self):
        audio = AudioFile(status=AudioFile.Status.DONE, language='ru')
        audio.file.save('s.wav', ContentFile(b'x'))
        Segment.objects.create(audio=audio, index=1, start=2.0, end=3.5, text='Второй')
        Segment.objects.create(audio=audio, index=0, start=0.0, end=2.0, text='Первый')

        data = self.client.get(reverse('player:subtitles', args=[audio.pk])).json()

        self.assertEqual(data['status'], 'done')
        self.assertEqual([s['text'] for s in data['segments']], ['Первый', 'Второй'])
        self.assertEqual(data['segments'][1], {'index': 1, 'start': 2.0, 'end': 3.5, 'text': 'Второй'})

    def test_title_defaults_to_file_name(self):
        audio = AudioFile()
        audio.file.save('my-call.wav', ContentFile(b'x'))
        self.assertEqual(audio.title, 'my-call')


class CallNameTests(TestCase):
    def test_parse(self):
        info = parse_call_name('out-375290000101-304-20260926-123305-1790415184.21439.wav')
        self.assertEqual(info['phone'], '375290000101')
        self.assertEqual(info['operator'], '304')
        self.assertEqual(info['call_id'], '1790415184.21439')
        started = timezone.localtime(info['call_started_at'])
        self.assertEqual(started.strftime('%Y-%m-%d %H:%M:%S'), '2026-09-26 12:33:05')

    def test_copy_suffix_and_path_are_ignored(self):
        info = parse_call_name('audio/out-375250000102-304-20260926-121304-1790413982.21348 (1).wav')
        self.assertEqual(info['call_id'], '1790413982.21348')

    def test_storage_suffix_is_ignored(self):
        info = parse_call_name('audio/out-375290000101-304-20260926-123305-1790415184_kRlAzSK.21439.wav')
        self.assertEqual(info['phone'], '375290000101')
        self.assertEqual(info['call_id'], '1790415184.21439')

    def test_other_names(self):
        self.assertIsNone(parse_call_name('interview.wav'))
        self.assertIsNone(parse_call_name('out-375-304-20261399-123305-1.2.wav'))  # bad date

    @override_settings(MEDIA_ROOT=MEDIA_ROOT)
    def test_filled_on_save(self):
        audio = AudioFile()
        audio.file.save('out-375290000101-304-20260926-123305-1790415184.21439.wav', ContentFile(b'x'))
        audio.refresh_from_db()
        self.assertEqual(audio.phone_display, '+375 (29) 000-01-01')
        self.assertEqual(audio.operator, '304')


@override_settings(MEDIA_ROOT=MEDIA_ROOT)
class DownloadSubtitlesTests(TestCase):
    def setUp(self):
        self.audio = AudioFile(status=AudioFile.Status.DONE)
        self.audio.file.save('out-375290000101-304-20260926-123305-1790415184.21439.wav', ContentFile(b'x'))
        Segment.objects.create(audio=self.audio, index=0, start=0.0, end=2.5, text='Добрый день!')
        Segment.objects.create(audio=self.audio, index=1, start=3661.25, end=3662.0, text='До свидания.')

    def download(self, fmt):
        return self.client.get(reverse('player:download_subtitles', args=[self.audio.pk, fmt]))

    def test_srt(self):
        r = self.download('srt')
        self.assertEqual(r.content.decode(), (
            '1\n00:00:00,000 --> 00:00:02,500\nДобрый день!\n\n'
            '2\n01:01:01,250 --> 01:01:02,000\nДо свидания.\n'
        ))
        self.assertIn('attachment', r['Content-Disposition'])
        self.assertIn('375290000101_304_2026-09-26_12-33-05.srt', r['Content-Disposition'])

    def test_vtt(self):
        body = self.download('vtt').content.decode()
        self.assertTrue(body.startswith('WEBVTT\n\n00:00:00.000 --> 00:00:02.500\nДобрый день!\n'))

    def test_txt_has_call_details(self):
        body = self.download('txt').content.decode()
        self.assertIn('Телефон: +375 (29) 000-01-01', body)
        self.assertIn('Оператор: 304', body)
        self.assertIn('Дата и время: 26.09.2026 12:33:05', body)
        self.assertIn('[01:01:01] До свидания.', body)

    def test_unknown_format(self):
        self.assertEqual(self.download('docx').status_code, 404)

    def test_copy_to_clipboard_uses_txt_export(self):
        r = self.client.get(self.audio.get_absolute_url())
        txt_url = reverse('player:download_subtitles', args=[self.audio.pk, 'txt'])
        self.assertContains(r, f'id="copy-subtitles" data-url="{txt_url}"')


class UploadTests(TestCase):
    def test_rejects_non_wav(self):
        r = self.client.post(reverse('player:upload'), {
            'files': ContentFile(b'x', name='song.mp3'),
        }, follow=True)
        self.assertContains(r, 'Поддерживаются только файлы .wav')
        self.assertFalse(AudioFile.objects.exists())

    @override_settings(MEDIA_ROOT=MEDIA_ROOT)
    def test_pbx_name_gives_call_details(self):
        r = self.client.post(reverse('player:upload'), {'files': ContentFile(make_wav(), name=CALL_A)})
        audio = AudioFile.objects.get()
        self.assertRedirects(r, audio.get_absolute_url())
        self.assertEqual((audio.phone, audio.operator), ('375290000103', '308'))
        self.assertEqual(timezone.localtime(audio.call_started_at).date(), date(2026, 9, 25))
        self.assertEqual(audio.status, AudioFile.Status.PENDING)  # queued for the worker
        self.assertAlmostEqual(audio.duration, 1.0)

    @override_settings(MEDIA_ROOT=MEDIA_ROOT)
    def test_other_name_requires_date(self):
        r = self.client.post(reverse('player:upload'), {'files': ContentFile(make_wav(), name='Интервью.wav')})
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'укажите дату и время звонка')
        self.assertContains(r, '<details id="upload" open>', html=False)
        self.assertFalse(AudioFile.objects.exists())

    @override_settings(MEDIA_ROOT=MEDIA_ROOT)
    def test_other_name_with_entered_details(self):
        self.client.post(reverse('player:upload'), {
            'files': ContentFile(make_wav(), name='Интервью.wav'),
            'call_started_at': '2026-09-27T10:15', 'operator': '308', 'phone': '+375 (29) 000-00-99',
        })
        audio = AudioFile.objects.get()
        self.assertEqual(timezone.localtime(audio.call_started_at).strftime('%Y-%m-%d %H:%M'), '2026-09-27 10:15')
        self.assertEqual((audio.operator, audio.phone), ('308', '375290000099'))
        # The call now shows up under its day.
        r = self.client.get(reverse('player:index') + '?day=2026-09-27')
        self.assertContains(r, audio.get_absolute_url())

    @override_settings(MEDIA_ROOT=MEDIA_ROOT)
    def test_many_files_at_once(self):
        make_call(CALL_C)  # already uploaded earlier
        files = [
            ContentFile(make_wav(), name=CALL_A),
            ContentFile(make_wav(), name=CALL_B),
            ContentFile(make_wav(), name=CALL_C),
            ContentFile(make_wav(), name=CALL_A),  # twice in the same batch
            ContentFile(make_wav(), name='Интервью 1.wav'),
            ContentFile(make_wav(), name='Интервью 2.wav'),
        ]
        r = self.client.post(reverse('player:upload'), {
            'files': files, 'call_started_at': '2026-09-27T10:15', 'operator': '308',
        }, follow=True)
        self.assertEqual(AudioFile.objects.count(), 5)  # 1 old + A, B + two interviews
        self.assertContains(r, 'Загружено файлов: 4')
        self.assertContains(r, f'пропущено 2: {CALL_C}, {CALL_A}')
        self.assertEqual(r.redirect_chain[-1][0], reverse('player:index') + '?day=2026-09-25')
        interview = AudioFile.objects.get(title='Интервью 1')
        self.assertEqual(interview.operator, '308')
        self.assertEqual(timezone.localtime(interview.call_started_at).strftime('%d.%m %H:%M'), '27.09 10:15')
        self.assertEqual(interview.status, AudioFile.Status.PENDING)
        self.assertAlmostEqual(interview.duration, 1.0)

    @override_settings(MEDIA_ROOT=MEDIA_ROOT)
    def test_batch_of_duplicates_opens_their_day(self):
        make_call(CALL_A)
        make_call(CALL_B)
        files = [ContentFile(make_wav(), name=CALL_A), ContentFile(make_wav(), name=CALL_B)]
        r = self.client.post(reverse('player:upload'), {'files': files})
        self.assertRedirects(r, reverse('player:index') + '?day=2026-09-25')
        self.assertEqual(AudioFile.objects.count(), 2)

    @override_settings(MEDIA_ROOT=MEDIA_ROOT)
    def test_batch_other_names_need_date_and_single_phone(self):
        files = [ContentFile(make_wav(), name='a.wav'), ContentFile(make_wav(), name='b.wav')]
        r = self.client.post(reverse('player:upload'), {'files': files, 'phone': '375290000001'})
        self.assertContains(r, 'укажите дату и время звонка')
        self.assertContains(r, 'a.wav, b.wav')
        self.assertContains(r, 'Телефон можно указать, только когда файл не из АТС один.')
        self.assertFalse(AudioFile.objects.exists())

    @override_settings(MEDIA_ROOT=MEDIA_ROOT)
    def test_batch_rejects_non_wav(self):
        files = [ContentFile(make_wav(), name=CALL_A), ContentFile(b'x', name='song.mp3')]
        r = self.client.post(reverse('player:upload'), {'files': files})
        self.assertContains(r, 'Поддерживаются только файлы .wav: song.mp3')
        self.assertFalse(AudioFile.objects.exists())

    @override_settings(MEDIA_ROOT=MEDIA_ROOT)
    def test_same_call_twice_opens_existing(self):
        existing = make_call(CALL_A)
        r = self.client.post(reverse('player:upload'), {'files': ContentFile(make_wav(), name=CALL_A)})
        self.assertRedirects(r, existing.get_absolute_url())
        self.assertEqual(AudioFile.objects.count(), 1)


def make_wav(seconds=1.0, rate=8000) -> bytes:
    buf = io.BytesIO()
    with wave.open(buf, 'wb') as w:
        w.setnchannels(1)
        w.setsampwidth(2)
        w.setframerate(rate)
        w.writeframes(b'\x00\x00' * int(seconds * rate))
    return buf.getvalue()


def make_call(name, duration=120.0, **fields):
    audio = AudioFile(duration=duration, **fields)
    audio.file.save(name, ContentFile(b'RIFF'))
    return audio


CALL_A = 'out-375290000103-308-20260925-093215-1790000001.100.wav'
CALL_B = 'out-375290000104-308-20260925-095111-1790000002.101.wav'
CALL_C = 'out-375290000105-306-20260925-091817-1790000003.102.wav'


@override_settings(MEDIA_ROOT=MEDIA_ROOT)
class ReportTests(TestCase):
    def setUp(self):
        Interviewer.objects.create(extension='308', name='Марина')
        Interviewer.objects.create(extension='306', name='Елена')
        self.a = make_call(CALL_A, duration=122.4)
        self.b = make_call(CALL_B, duration=171)
        self.c = make_call(CALL_C, duration=123)
        CallReview.objects.create(
            audio=self.a, review_date=date(2026, 9, 26), interviewer='Марина 308',
            city='Брест', stations='Не слушает', result='ок', controller='ИБ', completed=True,
        )
        CallReview.objects.create(
            audio=self.c, review_date=date(2026, 9, 26), interviewer='Елена 306', city='Минск',
            has_errors=True, error_comment='неверно указан возраст', result='ошибка',
            controller='ИБ', fixed='исправ ИБ', completed=True,
        )
        # Another day must not leak into the report.
        make_call('out-375290000000-308-20260926-100000-1790000009.109.wav')

    def load(self):
        r = self.client.get(reverse('player:report', args=['2026-09-25']))
        self.assertEqual(r.status_code, 200)
        self.assertIn('2026', r['Content-Disposition'])
        return openpyxl.load_workbook(io.BytesIO(r.content))['Лист1']

    def test_layout_matches_template(self):
        ws = self.load()
        headers = [ws.cell(row=1, column=i).value for i in range(1, 17)]
        self.assertEqual(headers, [h for h, _ in report.COLUMNS])
        self.assertEqual(ws['D1'].value, 'Дата опроса  Время')  # two spaces, as in the template
        self.assertTrue(ws['B1'].font.b)
        self.assertEqual(ws.freeze_panes, 'D2')
        self.assertEqual(ws.auto_filter.ref, 'A1:O4')
        self.assertAlmostEqual(ws.column_dimensions['K'].width, 29.3)
        self.assertEqual(ws['B2'].number_format, 'dd.mm.yyyy')
        self.assertEqual(ws['C2'].number_format, '0')  # number, shown in full
        self.assertEqual(ws['D2'].number_format, 'yyyy-mm-dd h:mm:ss')

    def test_invalid_dates(self):
        self.assertEqual(self.client.get('/report/2026-13-45.xlsx').status_code, 404)
        self.assertEqual(self.client.get(reverse('player:index') + '?day=2026-02-30').status_code, 200)
        r = self.client.post(reverse('player:sync_start'), {'day': '2026-02-30'}, follow=True)
        self.assertContains(r, 'Укажите день.')

    def test_styles_match_template(self):
        ws = self.load()
        self.assertEqual(ws.row_dimensions[1].height, 68)
        for col in 'BCDEN':
            header = ws[f'{col}1']
            self.assertEqual((header.alignment.horizontal, header.alignment.vertical), ('center', 'center'))
            self.assertEqual(header.border.bottom.style, 'thin')
        self.assertIsNone(ws['O1'].border.bottom.style)  # «Контроль» has no frame in the template
        self.assertEqual((ws['C1'].font.sz, ws['G1'].font.sz), (12, 11))
        self.assertEqual((ws['C2'].font.name, ws['C2'].alignment.horizontal), ('Arial', 'right'))
        self.assertEqual((ws['K2'].font.name, ws['K2'].alignment.wrap_text), ('Calibri', True))

    def test_empty_fields_filled_from_ai_suggestion(self):
        for audio, city, stations in [(self.a, 'Минск', 'Русское'), (self.b, 'Гродно', 'Супер FM')]:
            CallAnalysis.objects.create(audio=audio, status=CallAnalysis.Status.DONE, city=city, stations=stations)
        ws = self.load()
        rows = {ws[f'C{r}'].value: r for r in (2, 3, 4)}
        reviewed, unreviewed = rows[375290000103], rows[375290000104]
        # The controller's own value wins; only the empty field takes the suggestion.
        self.assertEqual(ws[f'G{reviewed}'].value, 'Брест')
        self.assertFalse(ws[f'G{reviewed}'].font.i)
        self.assertEqual(ws[f'I{reviewed}'].value, 'Не слушает')
        # Not reviewed: suggestions in grey italic with a note.
        self.assertEqual((ws[f'G{unreviewed}'].value, ws[f'I{unreviewed}'].value), ('Гродно', 'Супер FM'))
        self.assertTrue(ws[f'G{unreviewed}'].font.i)
        self.assertIn('Подсказка ИИ', ws[f'I{unreviewed}'].comment.text)
        self.assertIsNone(ws[f'H{unreviewed}'].value)  # empty suggestion stays empty

    def test_highlighting_like_controllers(self):
        CallReview.objects.create(audio=self.b, interviewer='Марина 308', result='ок',
                                  error_comment='Рекомендация: не подсказывать станцию')
        ws = self.load()
        # Row 2: Елена, «ошибка» -> whole row orange; row 4: «ок» with a comment -> green comment.
        self.assertEqual({ws[f'{c}2'].fill.fgColor.rgb for c in 'ABCDEFGHIJKLMNOP'}, {'FFFFC000'})
        self.assertEqual(ws['K4'].fill.fgColor.rgb, 'FF00FF00')
        self.assertEqual(ws['G4'].fill.fgColor.rgb, 'FFFFFFFF')
        self.assertEqual(ws['K3'].fill.fgColor.rgb, 'FFFFFFFF')  # «ок» without a comment

    def test_long_comment_makes_row_taller(self):
        CallReview.objects.filter(audio=self.c).update(error_comment='очень длинный комментарий контролёра ' * 4)
        ws = self.load()
        heights = [ws.row_dimensions[i].height for i in (2, 3, 4)]
        self.assertEqual(heights[1:], [16, 16])
        self.assertGreater(heights[0], 16)  # Елена's row with the long comment

    def test_rows(self):
        ws = self.load()
        rows = [[c.value for c in row] for row in ws.iter_rows(min_row=2, max_col=16)]
        self.assertEqual(len(rows), 3)
        # Sorted by interviewer, then time: Елена first, then Марина's two calls.
        self.assertEqual([r[4] for r in rows], ['Елена 306', 'Марина 308', 'Марина 308'])
        elena, marina_reviewed, marina_open = rows
        self.assertEqual(elena[2], 375290000105)
        self.assertEqual(elena[3], datetime(2026, 9, 25, 9, 18, 17))
        self.assertEqual(elena[9], 'есть')
        self.assertEqual(elena[10], 'неверно указан возраст')
        self.assertEqual(elena[11], 'ошибка')
        self.assertEqual(elena[15], 'исправ ИБ')
        self.assertEqual(marina_reviewed[1], datetime(2026, 9, 26))
        self.assertEqual(marina_reviewed[6], 'Брест')
        self.assertEqual(marina_reviewed[13], 122)
        self.assertEqual(marina_reviewed[14], 'ИБ')
        # Not reviewed yet: auto fields only, interviewer from the directory.
        self.assertIsNone(marina_open[1])
        self.assertEqual(marina_open[13], 171)
        self.assertIsNone(marina_open[11])


@override_settings(MEDIA_ROOT=MEDIA_ROOT)
class ReviewTests(TestCase):
    def setUp(self):
        AppSettings.objects.create(pk=1, controller='ИБ')
        Interviewer.objects.create(extension='308', name='Марина')
        self.a = make_call(CALL_A)
        self.b = make_call(CALL_B)

    def test_form_prefilled(self):
        r = self.client.get(self.a.get_absolute_url())
        self.assertContains(r, 'value="Марина 308"')
        self.assertContains(r, 'value="ИБ"')

    def test_autosave_then_complete(self):
        url = reverse('player:review_save', args=[self.a.pk])
        data = {'interviewer': 'Марина 308', 'review_date': '2026-09-26', 'city': 'Минск', 'result': 'ок'}
        r = self.client.post(url, data)
        self.assertEqual(r.json()['completed'], False)
        self.assertEqual(CallReview.objects.get(audio=self.a).city, 'Минск')

        r = self.client.post(url, {**data, 'complete': '1'})
        self.assertTrue(r.json()['completed'])
        self.assertEqual(r.json()['next_url'], self.b.get_absolute_url())
        self.assertEqual(CallReview.objects.count(), 1)

    def test_invalid_result(self):
        r = self.client.post(reverse('player:review_save', args=[self.a.pk]), {'result': 'отлично', 'review_date': '2026-09-26'})
        self.assertEqual(r.status_code, 400)

    def test_default_interviewer_remembers_last_choice(self):
        CallReview.objects.create(audio=self.a, interviewer='Маргарита 308')
        self.assertEqual(self.b.default_interviewer(), 'Маргарита 308')


class FakePbx:
    def __init__(self, recordings, body=None):
        self.recordings = recordings
        self.body = body or make_wav()
        self.downloaded = []

    def list_recordings(self, day, min_duration=0, max_duration=0):
        self.limits = (min_duration, max_duration)
        return self.recordings

    def download(self, rec, dest, debug_dir=None):
        if rec.filename in getattr(self, 'broken', ()):
            raise PbxError(f'Не удалось скачать запись {rec.filename} — download_audio: HTTP 200, text/html')
        self.downloaded.append(rec.filename)
        Path(dest).write_bytes(self.body)


def rec(filename, billsec=100, uniqueid='u'):
    return Recording(uniqueid=uniqueid, filename=filename, calldate=datetime(2026, 9, 25),
                     src='308', dst='', billsec=billsec, disposition='ANSWERED')


@override_settings(MEDIA_ROOT=MEDIA_ROOT)
class SyncTests(TestCase):
    def setUp(self):
        AppSettings.objects.create(pk=1, min_duration=30)
        Interviewer.objects.create(extension='308', name='Марина')

    def test_downloads_new_calls_and_skips_known(self):
        make_call(CALL_A)  # already downloaded before
        fake = FakePbx([
            rec(CALL_A),
            rec(CALL_B),
            rec(CALL_C),                                    # extension 306 is not in the list
            rec('out-375290000001-308-20260925-120000-1790000010.110.wav', billsec=5),  # too short
        ])
        sync = PbxSync.objects.create(day=date(2026, 9, 25))
        run_sync(sync, client=fake)

        sync.refresh_from_db()
        self.assertEqual(sync.status, PbxSync.Status.DONE)
        self.assertEqual((sync.found, sync.downloaded, sync.skipped), (2, 1, 1))
        self.assertEqual(fake.downloaded, [CALL_B])
        self.assertEqual(fake.limits, (30, 900))  # sent to the PBX search
        new = AudioFile.objects.get(call_id='1790000002.101')
        self.assertEqual(new.status, AudioFile.Status.PENDING)
        self.assertEqual(new.sync, sync)
        self.assertAlmostEqual(new.duration, 1.0)

    def test_failed_download_does_not_stop_the_day(self):
        fake = FakePbx([rec(CALL_A, uniqueid='a'), rec(CALL_B, uniqueid='b')])
        fake.broken = {CALL_A}
        sync = PbxSync.objects.create(day=date(2026, 9, 25))
        run_sync(sync, client=fake)
        sync.refresh_from_db()
        self.assertEqual((sync.status, sync.found, sync.downloaded, sync.failed), ('done', 2, 1, 1))
        self.assertIn('Не скачано 1 из 2', sync.error)
        self.assertIn('text/html', sync.error)
        self.assertEqual(fake.downloaded, [CALL_B])

    def test_all_downloads_failed_is_an_error(self):
        fake = FakePbx([rec(CALL_A, uniqueid='a')])
        fake.broken = {CALL_A}
        sync = PbxSync.objects.create(day=date(2026, 9, 25))
        run_sync(sync, client=fake)
        sync.refresh_from_db()
        self.assertEqual((sync.status, sync.failed), ('error', 1))

    def test_pbx_error_is_reported(self):
        class Broken(FakePbx):
            def list_recordings(self, day, min_duration=0, max_duration=0):
                raise PbxError('Не удалось войти в АТС: неверный логин или пароль.')

        sync = PbxSync.objects.create(day=date(2026, 9, 25))
        run_sync(sync, client=Broken([]))
        sync.refresh_from_db()
        self.assertEqual(sync.status, PbxSync.Status.ERROR)
        self.assertIn('неверный логин', sync.error)


# Structure of the real «Отчёты CDR» result page (FreePBX at the call centre),
# with made-up phone numbers. Playback rows are never closed, as on the PBX.
def cdr_row(time, phone, ext, uniqueid, talk='01:46', shown='01:47', status='ANSWERED', recorded=True):
    filename = f'out-{phone}-{ext}-{time[:10].replace("-", "")}-{time[11:].replace(":", "")}-{uniqueid}.wav'
    recording = (
        f'''<td title="{filename}"><a href="#" onClick="javascript:cdr_play(131,'{uniqueid}'); return false;">'''
        f'''<img src="assets/cdr/images/cdr_sound.png" alt="Call recording" /></a>\n    '''
        f'''<a href="/admin/config.php?display=cdr&action=download_audio&cdr_file={uniqueid}">'''
        f'''<img src="assets/cdr/images/cdr_download.png" alt="Call recording" /></a></td>'''
        if recorded else '<td></td>'
    )
    return (
        f'''<tr id="playback-{uniqueid}" class="playback" style="display:none;"><td colspan="14">'''
        f'''<div id="jquery_jplayer_1" class="jp-jplayer"></div><div class="jp-no-solution">'''
        f'''<span>Update Required</span></div>  <tr class="record">\n'''
        f'''<td>{time}</td>{recording}'''
        f'''<td title="Идентификатор: {uniqueid}"><a href="/admin/config.php?display=cdr&action=cel_show&uid={uniqueid}" >{uniqueid}</a></td>'''
        f'''<td title="Канал: SIP/{ext}-0000631e">&quot;{ext}&quot; &lt;{ext}&gt;</td>'''
        f'''<td title="Канал: SIP/10017-0000631f">&lt;{ext}&gt;</td><td title="Входящий номер: "></td>'''
        f'''<td title="Приложение: Dial(SIP/10017/80{phone[3:]},300,TM(outdial))">Dial</td>'''
        f'''<td title="Канал: SIP/10017-0000631f Контекст: from-internal">{phone}</td>'''
        f'''<td title="AMA-флаг: DEFAULT">{status}</td><td title="Время разговора: {talk}">{shown}</td>'''
        f'''<td></td><td></td>    <td></td>\n    <td></td>\n  </tr>\n'''
    )


CDR_PAGE = (
    '<html><body><form method="post" action="config.php?display=cdr">'
    '<input name="startday" value="29"><input name="endday" value="29"></form>'
    '<table class="cdr">'
    + cdr_row('2026-09-29 14:01:00', '375290000201', '301', '1790679660.25374')
    + cdr_row('2026-09-29 14:10:05', '375290000202', '305', '1790680205.25380', talk='12:03', shown='12:04')
    + cdr_row('2026-09-29 14:20:00', '375290000203', '301', '1790680800.25390', status='NO ANSWER',
              talk='00:00', shown='00:00', recorded=False)
    + '</table></body></html>'
)


def reply(body: bytes, content_type='audio/x-wav', status=200):
    """A streamed requests response; iter_content yields the (already decoded) body."""
    response = mock.Mock(status_code=status, headers={'Content-Type': content_type})
    response.iter_content.side_effect = lambda size: [body[i:i + size] for i in range(0, len(body), size)] or [b'']
    response.raw = io.BytesIO(b'\x1f\x8b compressed bytes')  # what .raw would give for gzip
    return response


class PbxClientTests(TestCase):
    def test_parse_cdr_html(self):
        recs = parse_cdr_html(CDR_PAGE)
        self.assertEqual(len(recs), 2)  # the call without a recording is dropped
        first = recs[0]
        self.assertEqual(first.filename, 'out-375290000201-301-20260929-140100-1790679660.25374.wav')
        self.assertEqual(first.uniqueid, '1790679660.25374')
        self.assertEqual(first.calldate, datetime(2026, 9, 29, 14, 1, 0))
        self.assertEqual((first.src, first.dst, first.disposition), ('301', '375290000201', 'ANSWERED'))
        self.assertEqual(first.billsec, 106)  # «Время разговора: 01:46», not the shown 01:47
        self.assertEqual(recs[1].billsec, 723)
        info = parse_call_name(first.filename)
        self.assertEqual((info['call_id'], info['operator']), ('1790679660.25374', '301'))

    def test_empty_result(self):
        self.assertEqual(parse_cdr_html('<form><input name="startday"></form><table></table>'), [])

    def test_unexpected_page(self):
        with self.assertRaises(PbxError):
            parse_cdr_html('<html><body>Доступ запрещён</body></html>')

    def test_seconds(self):
        self.assertEqual([pbx._seconds(t) for t in ('Время разговора: 01:46', '1:02:03', '')], [106, 3723, 0])

    def test_search_form_as_browser_sends_it(self):
        client = FreePbxClient('http://192.168.3.80', 'admin', 'secret')
        client._logged_in = True
        page = mock.Mock(status_code=200, text=CDR_PAGE)
        with mock.patch.object(client.session, 'request', return_value=page) as request:
            recs = client.list_recordings(date(2026, 9, 29), 80, 900)
        self.assertEqual(len(recs), 2)
        method, url = request.call_args.args
        form = request.call_args.kwargs['data']
        self.assertEqual((method, url), ('POST', 'http://192.168.3.80/admin/config.php?display=cdr'))
        self.assertEqual(
            {k: form[k] for k in ('startday', 'startmonth', 'startyear', 'endday', 'endhour', 'endmin',
                                  'need_html', 'limit', 'dur_min', 'dur_max', 'disposition', 'group')},
            {'startday': '29', 'startmonth': '09', 'startyear': '2026', 'endday': '29', 'endhour': '23',
             'endmin': '59', 'need_html': 'true', 'limit': '1000', 'dur_min': '80', 'dur_max': '900',
             'disposition': 'all', 'group': 'day'},
        )
        self.assertEqual((form['dst'], form['dst_mod']), ('', 'begins_with'))

    def test_login_like_browser(self):
        client = FreePbxClient('http://192.168.3.80/', 'operator', 'p@ss')
        with mock.patch.object(client.session, 'request', return_value=mock.Mock(status_code=200, text=CDR_PAGE)) as request:
            client.login()
        self.assertEqual(request.call_args.args, ('POST', 'http://192.168.3.80/admin/config.php?display=cdr'))
        self.assertEqual(request.call_args.kwargs['data'], {'username': 'operator', 'password': 'p@ss'})
        self.assertEqual(request.call_args.kwargs['headers'], {
            'Origin': 'http://192.168.3.80', 'Referer': 'http://192.168.3.80/admin/config.php?display=cdr',
        })

    def test_wrong_password(self):
        client = FreePbxClient('http://192.168.3.80', 'operator', 'wrong')
        login_page = mock.Mock(status_code=200, text='<form id="loginform"><input name="password"></form>')
        with mock.patch.object(client.session, 'request', return_value=login_page), \
                self.assertRaisesRegex(PbxError, 'неверный логин или пароль'):
            client.login()

    def test_download_reads_decoded_body(self):
        client = FreePbxClient('http://192.168.3.80', 'admin', 'secret')
        client._logged_in = True
        with mock.patch.object(client.session, 'request', return_value=reply(make_wav())), \
                tempfile.TemporaryDirectory() as tmp:
            dest = Path(tmp) / CALL_A
            client.download(rec(CALL_A), dest)
            self.assertTrue(dest.read_bytes().startswith(b'RIFF'))  # not the gzip bytes of .raw

    def test_download_cuts_html_before_wav(self):
        client = FreePbxClient('http://192.168.3.80', 'admin', 'secret')
        client._logged_in = True
        body = b'<!DOCTYPE html><html>FreePBX header</html>\n' + make_wav()
        with mock.patch.object(client.session, 'request', return_value=reply(body, 'text/html')), \
                tempfile.TemporaryDirectory() as tmp:
            dest = Path(tmp) / CALL_A
            client.download(rec(CALL_A), dest)
            self.assertEqual(dest.read_bytes(), make_wav())

    def test_non_wav_reply_is_described_and_saved(self):
        client = FreePbxClient('http://192.168.3.80', 'admin', 'secret')
        client._logged_in = True
        page = reply('<html><body>Файл записи не найден</body></html>'.encode(), 'text/html; charset=UTF-8')
        with mock.patch.object(client.session, 'request', return_value=page), \
                tempfile.TemporaryDirectory() as tmp:
            debug = Path(tmp) / 'debug'
            with self.assertRaises(PbxError) as ctx:
                client.download(rec(CALL_A, uniqueid='1790679660.25374'), Path(tmp) / CALL_A, debug_dir=debug)
            message = str(ctx.exception)
            self.assertIn('download_audio: HTTP 200, text/html; charset=UTF-8', message)
            self.assertIn('Файл записи не найден', message)
            self.assertEqual(sorted(p.name for p in debug.iterdir()),
                             ['1790679660.25374-ajax.html', '1790679660.25374-download_audio.html'])

    def test_download_uses_confirmed_link_first(self):
        client = FreePbxClient('http://192.168.3.80', 'admin', 'secret')
        client._logged_in = True
        ok = reply(make_wav())
        with mock.patch.object(client.session, 'request', return_value=ok) as request, \
                tempfile.TemporaryDirectory() as tmp:
            client.download(rec(CALL_A, uniqueid='1790679660.25374'), Path(tmp) / CALL_A)
        self.assertEqual(
            request.call_args_list[0].args[1],
            'http://192.168.3.80/admin/config.php?display=cdr&action=download_audio&cdr_file=1790679660.25374',
        )

    def test_download_rejects_login_page(self):
        client = FreePbxClient('https://pbx.local', 'admin', 'secret')
        client._logged_in = True
        html = reply(b'<html>login</html>', 'text/html; charset=UTF-8')
        with mock.patch.object(client.session, 'request', return_value=html), \
                tempfile.TemporaryDirectory() as tmp:
            with self.assertRaises(PbxError):
                client.download(rec(CALL_A), Path(tmp) / CALL_A)
            self.assertEqual(list(Path(tmp).iterdir()), [])  # no .part leftovers

    def test_download_saves_wav(self):
        client = FreePbxClient('https://pbx.local', 'admin', 'secret')
        client._logged_in = True
        ok = reply(make_wav())
        with mock.patch.object(client.session, 'request', return_value=ok), \
                tempfile.TemporaryDirectory() as tmp:
            dest = Path(tmp) / CALL_A
            client.download(rec(CALL_A), dest)
            self.assertTrue(dest.read_bytes().startswith(b'RIFF'))

    def test_missing_credentials(self):
        with self.assertRaises(PbxError):
            FreePbxClient('', 'admin', '')


@override_settings(MEDIA_ROOT=MEDIA_ROOT)
class WorkerTests(TestCase):
    def test_sync_runs_before_transcription(self):
        audio = make_call(CALL_A)
        sync = PbxSync.objects.create(day=date(2026, 9, 25))
        with mock.patch('player.worker.run_sync') as run_sync_mock, \
                mock.patch('player.worker.transcribe') as transcribe_mock:
            self.assertTrue(worker.run_once())
            run_sync_mock.assert_called_once_with(sync)
            transcribe_mock.assert_not_called()

            PbxSync.objects.update(status=PbxSync.Status.DONE)
            self.assertTrue(worker.run_once())
            transcribe_mock.assert_called_once_with(audio)
        self.assertTrue(AppSettings.load().worker_alive)

    def test_idle(self):
        self.assertFalse(worker.run_once())

    def test_recover_interrupted(self):
        audio = make_call(CALL_A, status=AudioFile.Status.PROCESSING)
        worker.recover_interrupted()
        audio.refresh_from_db()
        self.assertEqual(audio.status, AudioFile.Status.PENDING)


class EngineTests(TestCase):
    def test_available_engines_does_not_import(self):
        # Importing both torch and CTranslate2 in one process segfaults.
        from .transcription import available_engines

        with mock.patch.dict('sys.modules'):
            for module in ('whisper', 'faster_whisper'):
                sys.modules.pop(module, None)
            available_engines()
            self.assertNotIn('whisper', sys.modules)
            self.assertNotIn('faster_whisper', sys.modules)


class FakeKeyring:
    """In-memory stand-in for the OS credential store."""

    def __init__(self, keep=True):
        self.store = {}
        self.keep = keep

    def set_password(self, service, username, password):
        if self.keep:
            self.store[service, username] = password

    def get_password(self, service, username):
        return self.store.get((service, username))

    def get_keyring(self):
        return self


class SettingsPageTests(TestCase):
    DATA = {'pbx_url': '192.168.3.80', 'pbx_username': 'operator', 'pbx_password': 'secret',
            'controller': 'ИБ', 'min_duration': 80, 'max_duration': 900, 'interviewers': ''}

    def post(self, keyring, **changes):
        with mock.patch('player.credentials.keyring', keyring):
            return self.client.post(reverse('player:settings'), {**self.DATA, **changes}, follow=True)

    def test_password_saved_and_shown_as_saved(self):
        keyring = FakeKeyring()
        r = self.post(keyring)
        self.assertEqual(keyring.store, {('SyncVoice PBX', 'operator'): 'secret'})
        self.assertContains(r, 'Настройки сохранены. Пароль АТС сохранён.')
        self.assertContains(r, 'Пароль для логина «operator» сохранён')
        self.assertNotContains(r, 'secret')  # never rendered back

    def test_address_without_scheme_is_http(self):
        self.post(FakeKeyring())
        self.assertEqual(AppSettings.load().pbx_url, 'http://192.168.3.80')

    def test_store_that_does_not_keep_the_password(self):
        r = self.post(FakeKeyring(keep=False))
        self.assertContains(r, 'Не удалось сохранить пароль')
        self.assertContains(r, 'Настройки не сохранены')
        self.assertEqual(AppSettings.load().pbx_username, '')

    def test_renamed_login_keeps_password(self):
        keyring = FakeKeyring()
        self.post(keyring)
        r = self.post(keyring, pbx_username='operator2', pbx_password='')
        self.assertEqual(keyring.store[('SyncVoice PBX', 'operator2')], 'secret')
        self.assertContains(r, 'Пароль для логина «operator2» сохранён')

    def test_empty_password_keeps_the_saved_one(self):
        keyring = FakeKeyring()
        self.post(keyring)
        self.post(keyring, pbx_password='', controller='АБ')
        self.assertEqual(keyring.store[('SyncVoice PBX', 'operator')], 'secret')
        self.assertEqual(AppSettings.load().controller, 'АБ')


class SettingsFormTests(TestCase):
    def test_interviewers_parsed_and_replaced(self):
        Interviewer.objects.create(extension='301', name='Эдуард')
        form = SettingsForm({
            'min_duration': 30, 'max_duration': 900,
            'interviewers': '301 Виолетта\n\n308 Марина\n308 Маргарита\n',
        }, instance=AppSettings.load())
        self.assertTrue(form.is_valid(), form.errors)
        form.save()
        form.save_interviewers()
        self.assertEqual(
            sorted(map(str, Interviewer.objects.all())),
            ['Виолетта 301', 'Маргарита 308', 'Марина 308'],
        )

    def test_bad_line(self):
        form = SettingsForm({'min_duration': 30, 'max_duration': 900, 'interviewers': 'Марина'}, instance=AppSettings.load())
        self.assertFalse(form.is_valid())
        self.assertIn('Строка 1', str(form.errors))


class FakeGemini:
    """Stands in for google.genai.Client: returns `result` or raises it."""

    def __init__(self, result):
        self.result = result
        self.calls = []
        self.models = self

    def generate_content(self, **kwargs):
        self.calls.append(kwargs)
        if isinstance(self.result, Exception):
            raise self.result
        return SimpleNamespace(
            parsed=self.result, text='',
            usage_metadata=SimpleNamespace(prompt_token_count=900, candidates_token_count=60, thoughts_token_count=40),
        )


SUGGESTION = SuggestedFields(city='Минск', listen_city='', stations='Русское', notes='Слушала вчера на работе.')


@override_settings(MEDIA_ROOT=MEDIA_ROOT, GEMINI_API_KEY='test-key', GEMINI_MODEL='gemini-test')
class AnalysisTests(TestCase):
    def setUp(self):
        RadioStation.objects.create(name='Радио Юнистар', report_name='Юнистар', aliases='Unistar Radio',
                                    frequencies={'Минск': '99.5'})
        self.audio = make_call(CALL_A, status=AudioFile.Status.DONE)
        Segment.objects.create(audio=self.audio, index=0, start=16, end=18, text='Вы в каком городе проживаете? Минск.')
        Segment.objects.create(audio=self.audio, index=1, start=75, end=80, text='Вчера слушала Русское радио.')

    def analyse(self, result):
        analysis.queue(self.audio)
        item = self.audio.analysis
        fake = FakeGemini(result)
        analysis.run_analysis(item, client=fake)
        item.refresh_from_db()
        return item, fake

    def test_fields_saved(self):
        item, fake = self.analyse(SUGGESTION)
        self.assertEqual(item.status, CallAnalysis.Status.DONE)
        self.assertEqual((item.city, item.listen_city, item.stations), ('Минск', '', 'Русское'))
        self.assertEqual((item.input_tokens, item.output_tokens), (900, 100))
        self.assertEqual(item.model_name, 'gemini-test')
        request = fake.calls[0]
        self.assertEqual(request['model'], 'gemini-test')
        self.assertIn('[00:16] Вы в каком городе проживаете? Минск.', request['contents'])
        self.assertIn('- «Юнистар»; также: Радио Юнистар, Unistar Radio; частоты: Минск 99.5',
                      request['config'].system_instruction)  # station directory

    def test_only_transcript_is_sent(self):
        _, fake = self.analyse(SUGGESTION)
        sent = fake.calls[0]['contents'] + fake.calls[0]['config'].system_instruction
        self.assertNotIn(self.audio.phone, sent)
        self.assertNotIn('308', sent)

    def test_rejected_request_is_an_error(self):
        item, _ = self.analyse(genai_errors.ClientError(400, {'error': {'code': 400, 'message': 'API key not valid'}}))
        self.assertEqual(item.status, CallAnalysis.Status.ERROR)
        self.assertIn('API key not valid', item.error)

    def test_temporary_failures_are_retried(self):
        for exc in [
            genai_errors.ClientError(429, {'error': {'code': 429, 'message': 'Resource exhausted'}}),
            genai_errors.ServerError(503, {'error': {'code': 503, 'message': 'Overloaded'}}),
            httpx.ConnectError('proxy refused'),
        ]:
            with self.subTest(exc=type(exc).__name__), self.assertRaises(analysis.RetryLater):
                self.analyse(exc)
            self.assertEqual(self.audio.analysis.status, CallAnalysis.Status.PENDING)

    def test_review_form_prefilled_from_suggestion(self):
        self.analyse(SUGGESTION)
        r = self.client.get(self.audio.get_absolute_url())
        self.assertContains(r, 'name="city" value="Минск"', html=False)
        self.assertContains(r, 'value="Русское"')
        self.assertContains(r, 'id="ai-panel"')

    def test_existing_review_is_not_overwritten(self):
        CallReview.objects.create(audio=self.audio, city='Брест')
        self.analyse(SUGGESTION)
        r = self.client.get(self.audio.get_absolute_url())
        self.assertContains(r, 'name="city" value="Брест"', html=False)

    def test_run_button_queues(self):
        r = self.client.post(reverse('player:analysis', args=[self.audio.pk]))
        self.assertEqual(r.json()['analysis']['status'], 'pending')

    @override_settings(GEMINI_API_KEY='')
    def test_disabled_without_key(self):
        r = self.client.post(reverse('player:analysis', args=[self.audio.pk]))
        self.assertEqual(r.status_code, 400)
        self.assertNotContains(self.client.get(self.audio.get_absolute_url()), 'id="ai-panel"')

    def test_worker_queues_after_transcription_and_analyses(self):
        pending = make_call(CALL_B)
        def fake_transcribe(audio):
            Segment.objects.create(audio=audio, index=0, start=0, end=2, text='Минск.')
            AudioFile.objects.filter(pk=audio.pk).update(status='done')

        with mock.patch('player.worker.transcribe', side_effect=fake_transcribe):
            self.assertTrue(worker.run_once())
        self.assertEqual(pending.analysis.status, CallAnalysis.Status.PENDING)
        with mock.patch('player.analysis.make_client', return_value=FakeGemini(SUGGESTION)):
            self.assertTrue(worker.run_once())
        pending.analysis.refresh_from_db()
        self.assertEqual(pending.analysis.status, CallAnalysis.Status.DONE)

    def test_worker_backs_off_on_temporary_failure(self):
        analysis.queue(self.audio)
        failing = FakeGemini(httpx.ConnectError('down'))
        with mock.patch('player.analysis.make_client', return_value=failing):
            self.assertFalse(worker.run_analysis_step())
            self.assertFalse(worker.run_analysis_step())  # paused: no second request
        self.assertEqual(len(failing.calls), 1)
        self.assertIn('down', CallAnalysis.objects.get().error)
        worker._analysis_paused_until = 0


class ProxyTests(TestCase):
    @override_settings(GEMINI_PROXY_URL='http://10.0.0.5:3128', GEMINI_PROXY_USERNAME='admin',
                       GEMINI_PROXY_PASSWORD='p@ss:w/rd')
    def test_credentials_are_escaped(self):
        self.assertEqual(analysis.proxy_url(), 'http://admin:p%40ss%3Aw%2Frd@10.0.0.5:3128')

    @override_settings(GEMINI_PROXY_URL='')
    def test_no_proxy(self):
        self.assertIsNone(analysis.proxy_url())

    @override_settings(GEMINI_API_KEY='k', GEMINI_PROXY_URL='http://10.0.0.5:3128',
                       GEMINI_PROXY_USERNAME='', GEMINI_PROXY_PASSWORD='')
    def test_client_uses_proxy_only_for_gemini(self):
        with mock.patch('google.genai.Client') as client_cls:
            analysis.make_client()
        options = client_cls.call_args.kwargs['http_options']
        self.assertEqual(options.client_args, {'trust_env': False, 'proxy': 'http://10.0.0.5:3128'})
        self.assertEqual(client_cls.call_args.kwargs['api_key'], 'k')


def make_stations_xlsx() -> io.BytesIO:
    """A small workbook in the format of «Список РСТ с частотами.xlsx»."""
    wb = openpyxl.Workbook()
    master = wb.active
    master.title = 'без англ'
    master.append([None, None, 'Минск', 'Брест', 'Витебск', 'Гомель', 'Гродно', 'Могилев', None])
    master.append([1, 'Юмор FM', 93.7, 87.5, 96.2, 92.1, 89.9, 91.9, None])
    master.append([2, 'Культура', 102.9, 88.5, 99.3, 91.5, '95.0', 99.1, None])
    master.append([3, 'Супер FM', 104.6, 96.4, None, '91.0', None, None, 'Ранее Би-Эй'])
    master.append([4, 'Правда Радио Гомель', None, None, ' ', '99.0', None, None, None])
    minsk = wb.create_sheet('Минск')
    minsk.append([30, 'Humor FM', 'Юмор FM', '93,7'])
    minsk.append([2, 'Kultura', 'Культура', '102,9'])
    minsk.append([None, 'добавились ретро фм', None, None])  # comment row
    gomel = wb.create_sheet('Гомель')
    gomel.append([20, 'Pravda Radio', 'Правда Радио', '99,0'])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


class StationDirectoryTests(TestCase):
    def test_import(self):
        created, updated = stations.import_workbook(make_stations_xlsx())
        self.assertEqual((created, updated), (4, 0))
        humor = RadioStation.objects.get(name='Юмор FM')
        self.assertEqual(humor.report_name, 'Юмор')  # as controllers write it
        self.assertEqual(humor.alias_list(), ['Humor FM'])
        self.assertEqual(humor.frequencies['Гродно'], '89.9')
        culture = RadioStation.objects.get(name='Культура')
        self.assertEqual(culture.frequencies['Гродно'], '95.0')
        self.assertEqual(culture.report_name, '')
        self.assertEqual(RadioStation.objects.get(name='Правда Радио Гомель').alias_list(),
                         ['Правда Радио', 'Pravda Radio'])  # matched by similar name
        self.assertEqual(RadioStation.objects.get(name='Супер FM').note, 'Ранее Би-Эй')
        self.assertNotIn('Витебск', RadioStation.objects.get(name='Правда Радио Гомель').frequencies)

    def test_reimport_keeps_manual_edits(self):
        stations.import_workbook(make_stations_xlsx())
        humor = RadioStation.objects.get(name='Юмор FM')
        humor.report_name = 'Юмор ФМ'
        humor.aliases = 'Humor FM; Юморина'
        humor.active = False
        humor.frequencies = {}
        humor.save()
        self.assertEqual(stations.import_workbook(make_stations_xlsx()), (0, 4))
        humor.refresh_from_db()
        self.assertEqual((humor.report_name, humor.aliases, humor.active), ('Юмор ФМ', 'Humor FM; Юморина', False))
        self.assertEqual(humor.frequencies['Минск'], '93.7')  # frequencies come from the file

    def test_prompt_block_uses_report_names_and_skips_inactive(self):
        stations.import_workbook(make_stations_xlsx())
        RadioStation.objects.filter(name='Культура').update(active=False)
        block = stations.prompt_block()
        self.assertIn('- «Юмор»; также: Юмор FM, Humor FM; частоты: Минск 93.7, Брест 87.5', block)
        self.assertIn('примечание: Ранее Би-Эй', block)
        self.assertNotIn('Культура', block)
        self.assertEqual(stations.answer_names(), ['Правда Радио Гомель', 'Супер FM', 'Юмор'])

    def test_page_edit_and_add(self):
        stations.import_workbook(make_stations_xlsx())
        page = self.client.get(reverse('player:stations'))
        self.assertContains(page, 'value="Юмор FM"')
        formset = page.context['formset']
        data = {
            'form-TOTAL_FORMS': str(len(formset.forms)), 'form-INITIAL_FORMS': str(formset.initial_form_count()),
            'form-MIN_NUM_FORMS': '0', 'form-MAX_NUM_FORMS': '1000',
        }
        for i, form in enumerate(formset.forms):
            for name, field in form.fields.items():
                value = form.initial.get(name, field.initial)
                if isinstance(field, forms.BooleanField):
                    if value:
                        data[f'form-{i}-{name}'] = 'on'
                elif value is not None and name != 'DELETE':
                    data[f'form-{i}-{name}'] = value.pk if hasattr(value, 'pk') else value
        humor_i = next(i for i, f in enumerate(formset.forms) if f.instance.name == 'Юмор FM')
        data[f'form-{humor_i}-freq_Минск'] = '93,8'
        new_i = len(formset.forms) - 1
        data[f'form-{new_i}-name'] = 'Радио Новинка'
        data[f'form-{new_i}-freq_Брест'] = '101.1'
        data[f'form-{new_i}-active'] = 'on'
        r = self.client.post(reverse('player:stations'), data)
        self.assertRedirects(r, reverse('player:stations'))
        self.assertEqual(RadioStation.objects.get(name='Юмор FM').frequencies['Минск'], '93.8')
        self.assertEqual(RadioStation.objects.get(name='Радио Новинка').frequencies, {'Брест': '101.1'})

    def test_bad_frequency(self):
        form = forms_module.StationForm(data={'name': 'X', 'freq_Минск': 'сто три', 'active': 'on'})
        self.assertFalse(form.is_valid())
        self.assertIn('freq_Минск', form.errors)

    def test_import_via_page(self):
        upload = ContentFile(make_stations_xlsx().getvalue(), name='РСТ.xlsx')
        r = self.client.post(reverse('player:stations_import'), {'file': upload}, follow=True)
        self.assertContains(r, 'Импорт: добавлено 4, обновлено 0.')

    @override_settings(MEDIA_ROOT=MEDIA_ROOT)
    def test_review_form_suggests_report_names(self):
        stations.import_workbook(make_stations_xlsx())
        audio = make_call(CALL_A)
        self.assertContains(self.client.get(audio.get_absolute_url()), '<option value="Юмор">')


SURVEY = 'http://crm.local/answers?phone={phone}&local={phone_local}&d={date}&ru={date_ru}&t={time}&op={operator}&id={call_id}'


@override_settings(MEDIA_ROOT=MEDIA_ROOT)
class SurveyLinkTests(TestCase):
    def setUp(self):
        self.audio = make_call(CALL_A)  # 375290000103, ext 308, 2026-09-25 09:32:15

    def test_all_placeholders(self):
        AppSettings.objects.update_or_create(pk=1, defaults={'survey_url': SURVEY})
        self.assertEqual(
            self.audio.survey_url(),
            'http://crm.local/answers?phone=375290000103&local=290000103&d=2026-09-25'
            '&ru=25.09.2026&t=09%3A32&op=308&id=1790000001.100',
        )

    def test_no_template_or_phone(self):
        self.assertIsNone(self.audio.survey_url())
        AppSettings.objects.update_or_create(pk=1, defaults={'survey_url': SURVEY})
        self.audio.phone = ''
        self.assertIsNone(self.audio.survey_url())

    def test_unknown_placeholder_rejected_and_scheme_added(self):
        base = {'min_duration': 80, 'max_duration': 900}
        form = SettingsForm({**base, 'survey_url': 'crm.local/a?p={phone}'}, instance=AppSettings.load())
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data['survey_url'], 'http://crm.local/a?p={phone}')
        form = SettingsForm({**base, 'survey_url': 'http://crm.local/a?p={tel}'}, instance=AppSettings.load())
        self.assertFalse(form.is_valid())
        self.assertIn('Неизвестные подстановки: tel', str(form.errors))

    def test_survey_links_of_neighbouring_calls(self):
        AppSettings.objects.update_or_create(pk=1, defaults={'survey_url': 'http://crm.local/a?p={phone}'})
        other = make_call(CALL_B)  # the next call of the same day
        r = self.client.get(self.audio.get_absolute_url())
        self.assertContains(r, 'id="next-call" data-survey-url="http://crm.local/a?p=375290000104"')
        data = {'interviewer': '308', 'review_date': '2026-09-26', 'complete': '1'}
        saved = self.client.post(reverse('player:review_save', args=[self.audio.pk]), data).json()
        self.assertEqual(saved['next_url'], other.get_absolute_url())
        self.assertEqual(saved['next_survey_url'], 'http://crm.local/a?p=375290000104')

    def test_phone_in_header_copies_digits(self):
        r = self.client.get(self.audio.get_absolute_url())
        self.assertContains(r, 'id="phone-copy" data-copy="375290000103"')
        self.assertContains(r, 'callbar.js')

    def test_call_page_has_survey_pane(self):
        r = self.client.get(self.audio.get_absolute_url())
        self.assertNotContains(r, 'data-pane="survey"')
        AppSettings.objects.update_or_create(pk=1, defaults={'survey_url': 'http://crm.local/a?p={phone}'})
        r = self.client.get(self.audio.get_absolute_url())
        self.assertContains(r, 'data-pane="survey"')
        self.assertContains(r, 'data-url="http://crm.local/a?p=375290000103"')
        self.assertContains(r, 'id="survey-frame"')
        self.assertContains(r, 'data-pane="review"')  # the review form is still there


class SurveyCheckTests(TestCase):
    URL = 'http://crm.local/a?p=375290000103'

    def setUp(self):
        AppSettings.objects.update_or_create(pk=1, defaults={'survey_url': 'http://crm.local/a?p={phone}'})
        cache.clear()

    def check(self, headers=None, error=None, url=URL):
        reply = mock.Mock(headers=headers or {})
        with mock.patch('player.views.requests.Session.get', return_value=reply, side_effect=error) as get:
            r = self.client.get(reverse('player:survey_check'), {'url': url})
        return r, get

    def test_embeddable(self):
        r, _ = self.check()
        self.assertEqual(r.json(), {'ok': True, 'embeddable': True, 'reason': ''})

    def test_x_frame_options(self):
        r, _ = self.check({'X-Frame-Options': 'SAMEORIGIN'})
        self.assertFalse(r.json()['embeddable'])
        self.assertIn('SAMEORIGIN', r.json()['reason'])

    def test_frame_ancestors(self):
        r, _ = self.check({'Content-Security-Policy': "default-src 'self'; frame-ancestors 'self'"})
        self.assertFalse(r.json()['embeddable'])
        cache.clear()
        r, _ = self.check({'Content-Security-Policy': 'frame-ancestors http://127.0.0.1:8765'})
        self.assertTrue(r.json()['embeddable'])

    def test_unreachable(self):
        r, _ = self.check(error=requests.exceptions.ConnectionError('refused'))
        self.assertIsNone(r.json()['embeddable'])

    def test_other_hosts_refused(self):
        r, get = self.check(url='http://example.com/secret')
        self.assertEqual(r.status_code, 400)
        get.assert_not_called()

    def test_cached_per_host(self):
        self.check()
        _, get = self.check()
        get.assert_not_called()


CRM_LOGIN_PAGE = """<html><body><form method="post" action="/login">
<input type="email" name="email" value=""><input type="password" name="password">
<input type="hidden" name="_csrf_token" value="tok123"><input type="checkbox" name="_remember_me">
<button>Войти</button></form></body></html>"""
CRM_FIND = [  # the shape of the CRM's /admin/Reports/find reply, made-up numbers
    {'Id': 100, 'MembersSurveyID': '1', 'MemberID': '9', 'Phone': '375290000103', 'OPUserID': 17,
     'OPTime': '25-09-2025 10:00:00'},
    {'Id': 200, 'MembersSurveyID': '2', 'MemberID': '9', 'Phone': '375290000103', 'OPUserID': 7,
     'OPTime': '25-09-2026 09:34:01'},
    {'Id': 201, 'MembersSurveyID': '3', 'MemberID': '9', 'Phone': '375290000103', 'OPUserID': 8,
     'OPTime': '25-09-2026 16:00:00'},
]


def crm_reply(text='', status=200, url='http://crm.local/', json_data=None, headers=None, redirect_to=None):
    response = mock.Mock(status_code=status, text=text, url=url, headers=headers or {}, encoding='utf-8')
    response.content = text.encode()
    response.is_redirect = redirect_to is not None
    if redirect_to:
        response.headers = {'Location': redirect_to, **(headers or {})}
    response.json.side_effect = (lambda: json_data) if json_data is not None else ValueError
    return response


class CrmClientTests(TestCase):
    def client_with(self, *replies):
        client = crm.CrmClient('http://crm.local', 'sv@example.com', 'secret')
        patcher = mock.patch.object(client.session, 'request', side_effect=list(replies))
        self.request = patcher.start()
        self.addCleanup(patcher.stop)
        return client

    def test_login_submits_form_like_browser(self):
        client = self.client_with(
            crm_reply(CRM_LOGIN_PAGE, url='http://crm.local/login'),
            crm_reply('<h1>Админка</h1>', url='http://crm.local/admin'),
        )
        client.login()
        method, url = self.request.call_args.args
        self.assertEqual((method, url), ('POST', 'http://crm.local/login'))
        self.assertEqual(self.request.call_args.kwargs['data'],
                         {'email': 'sv@example.com', 'password': 'secret', '_csrf_token': 'tok123'})

    def test_wrong_password(self):
        client = self.client_with(
            crm_reply(CRM_LOGIN_PAGE, url='http://crm.local/login'),
            crm_reply(CRM_LOGIN_PAGE, url='http://crm.local/login'),
        )
        with self.assertRaisesRegex(crm.CrmError, 'неверный логин или пароль'):
            client.login()

    def test_find_and_pick(self):
        client = self.client_with(
            crm_reply(CRM_LOGIN_PAGE, url='http://crm.local/login'),
            crm_reply('ok', url='http://crm.local/admin'),
            crm_reply('[]', url='http://crm.local/admin/Reports/find', json_data=CRM_FIND),
        )
        surveys = client.find('375290000103')
        self.assertEqual([s.id for s in surveys], [100, 200, 201])
        self.assertEqual(self.request.call_args.kwargs['data'], {'PhoneN': '375290000103'})
        chosen = crm.pick_survey(surveys, datetime(2026, 9, 25, 9, 32, 15))
        self.assertEqual(chosen.id, 200)  # same day, closest to the call
        self.assertEqual(chosen.path, 'admin/Reports/update200')
        self.assertIsNone(crm.pick_survey(surveys, datetime(2026, 9, 27, 9, 0)))

    def test_expired_session_logs_in_again(self):
        client = self.client_with(
            crm_reply(CRM_LOGIN_PAGE, url='http://crm.local/login'),
            crm_reply('ok', url='http://crm.local/admin'),
            crm_reply('', status=302, redirect_to='/login'),  # session expired
            crm_reply(CRM_LOGIN_PAGE, url='http://crm.local/login'),
            crm_reply('ok', url='http://crm.local/admin'),
            crm_reply('<h1>Анкета</h1>', url='http://crm.local/admin/Reports/update200'),
        )
        response = client.request('GET', 'admin/Reports/update200', allow_redirects=False)
        self.assertIn('Анкета', response.text)


class CrmGatewayTests(TestCase):
    def setUp(self):
        self.client_ = mock.Mock(base_url='http://192.168.12.230/')

    def test_only_viewing(self):
        # A form submission (no XHR header) and an XHR to a write-like address are refused.
        for path, headers in [('/admin/Reports/update200', {}),
                              ('/admin/Reports/update200', {'X-Requested-With': 'XMLHttpRequest'}),
                              ('/admin/Reports/saveAnswers', {'X-Requested-With': 'XMLHttpRequest'})]:
            status, _, body = crm_proxy.proxy_response(self.client_, 'POST', path, b'a=1', headers)
            self.assertEqual(status, 405, path)
            self.assertIn('Только просмотр', body.decode())
        self.client_.request.assert_not_called()

    def test_delete_link_is_refused_even_as_get(self):
        # The CRM's survey page links «Удалить анкету» to GET /admin/Reports/delete<Id>.
        for path in ('/admin/Reports/delete200935', '/logout'):
            status, _, _ = crm_proxy.proxy_response(self.client_, 'GET', path)
            self.assertEqual(status, 405, path)
        self.client_.request.assert_not_called()

    def test_background_data_load_is_forwarded(self):
        self.client_.url.side_effect = lambda path: 'http://192.168.12.230/' + path
        self.client_.request.return_value = crm_reply('{"answers": []}', headers={'Content-Type': 'application/json'})
        status, _, body = crm_proxy.proxy_response(
            self.client_, 'POST', '/admin/Reports/answers', b'MembersSurveyID=235692',
            {'X-Requested-With': 'XMLHttpRequest', 'Content-Type': 'application/x-www-form-urlencoded; charset=UTF-8'})
        self.assertEqual((status, body), (200, b'{"answers": []}'))
        kwargs = self.client_.request.call_args.kwargs
        self.assertEqual(kwargs['data'], b'MembersSurveyID=235692')
        self.assertEqual(kwargs['headers']['X-Requested-With'], 'XMLHttpRequest')

    def test_refused_body_does_not_break_the_next_request(self):
        # The 501 «Unsupported method ('MembersSurveyID=…GET')» seen in the field.
        fake = mock.Mock(base_url='http://192.168.12.230/')
        fake.request.return_value = crm_reply('<h1>Анкета</h1>', headers={'Content-Type': 'text/html'})
        server = ThreadingHTTPServer(('127.0.0.1', 0), crm_proxy._Handler)
        threading.Thread(target=server.serve_forever, daemon=True).start()
        self.addCleanup(server.shutdown)
        with mock.patch('player.crm_proxy.get_client', return_value=fake):
            conn = http.client.HTTPConnection('127.0.0.1', server.server_address[1], timeout=5)
            conn.request('POST', '/admin/Reports/update200', body='MembersSurveyID=235692',
                         headers={'Content-Type': 'application/x-www-form-urlencoded'})
            first = conn.getresponse()
            first.read()
            self.assertEqual(first.status, 405)
            conn = http.client.HTTPConnection('127.0.0.1', server.server_address[1], timeout=5)
            conn.request('GET', '/admin/Reports/update200')
            second = conn.getresponse()
            self.assertEqual((second.status, second.read()), (200, '<h1>Анкета</h1>'.encode()))

    @override_settings(CRM_PROXY_PORT=8766)
    def test_rewrites_links_and_allows_framing(self):
        page = '<a href="http://192.168.12.230/admin/Reports">Назад</a><img src="//192.168.12.230/logo.png">'
        self.client_.request.return_value = crm_reply(page, headers={
            'Content-Type': 'text/html; charset=UTF-8', 'X-Frame-Options': 'DENY', 'Set-Cookie': 'PHPSESSID=x',
            'Content-Security-Policy': "default-src 'self'; frame-ancestors 'none'",
        })
        status, headers, body = crm_proxy.proxy_response(self.client_, 'GET', '/admin/Reports/update200')
        names = {name.lower(): value for name, value in headers}
        self.assertEqual(status, 200)
        self.assertNotIn('x-frame-options', names)
        self.assertNotIn('set-cookie', names)  # SyncVoice's CRM session never reaches the browser
        self.assertEqual(names['content-security-policy'], "default-src 'self'")
        self.assertEqual(body.decode(), '<a href="http://127.0.0.1:8766/admin/Reports">Назад</a>'
                                        '<img src="http://127.0.0.1:8766/logo.png">')

    @override_settings(CRM_PROXY_PORT=8766)
    def test_redirect_location_rewritten(self):
        self.client_.request.return_value = crm_reply(
            '', status=302, redirect_to='http://192.168.12.230/admin/Reports')
        status, headers, _ = crm_proxy.proxy_response(self.client_, 'GET', '/admin')
        self.assertEqual((status, dict(headers)['Location']), (302, 'http://127.0.0.1:8766/admin/Reports'))

    def test_crm_unreachable(self):
        self.client_.request.side_effect = crm.CrmError('CRM недоступна: timeout')
        status, _, body = crm_proxy.proxy_response(self.client_, 'GET', '/admin')
        self.assertEqual(status, 502)
        self.assertIn('timeout', body.decode())


@override_settings(MEDIA_ROOT=MEDIA_ROOT, CRM_PROXY_PORT=8766)
class CrmCallPageTests(TestCase):
    def setUp(self):
        AppSettings.objects.update_or_create(pk=1, defaults={'crm_url': 'http://crm.local', 'crm_username': 'sv'})
        self.audio = make_call(CALL_A)  # 375290000103, 2026-09-25 09:32:15
        fake = mock.Mock(base_url='http://crm.local/')
        fake.find.return_value = [crm.parse_survey(row) for row in CRM_FIND]
        fake.url.side_effect = lambda path: 'http://crm.local/' + path
        patcher = mock.patch('player.crm.get_client', return_value=fake)
        patcher.start()
        self.addCleanup(patcher.stop)

    def test_lookup_marks_this_calls_survey(self):
        data = self.client.get(reverse('player:crm_lookup', args=[self.audio.pk])).json()
        self.assertTrue(data['ok'])
        self.assertEqual(data['selected'], 200)
        self.assertEqual([s['id'] for s in data['surveys']], [201, 200, 100])  # newest first
        chosen = next(s for s in data['surveys'] if s['id'] == 200)
        self.assertEqual(chosen['view_url'], 'http://127.0.0.1:8766/admin/Reports/update200')
        self.assertEqual(chosen['direct_url'], 'http://crm.local/admin/Reports/update200')
        self.assertEqual(chosen['time'], '25.09.2026 09:34:01')

    def test_open_redirects_to_the_crm(self):
        r = self.client.get(reverse('player:crm_open', args=[self.audio.pk]))
        self.assertRedirects(r, 'http://crm.local/admin/Reports/update200', fetch_redirect_response=False)

    def test_call_page_in_crm_mode(self):
        r = self.client.get(self.audio.get_absolute_url())
        self.assertContains(r, f'data-lookup-url="{reverse("player:crm_lookup", args=[self.audio.pk])}"')
        self.assertContains(r, f'data-url="{reverse("player:crm_open", args=[self.audio.pk])}"')


class AudioDayFolderTests(TestCase):
    def test_folder_from_pbx_name_entered_date_or_today(self):
        self.assertEqual(audio_paths.audio_upload_to(AudioFile(), CALL_A), f'audio/2026/09/25/{CALL_A}')
        entered = AudioFile(call_started_at=timezone.make_aware(datetime(2026, 9, 27, 10, 15)))
        self.assertEqual(audio_paths.audio_upload_to(entered, 'Интервью.wav'), 'audio/2026/09/27/Интервью.wav')
        today = timezone.localdate()
        self.assertEqual(audio_paths.audio_upload_to(AudioFile(), 'x.wav'), f'audio/{today:%Y/%m/%d}/x.wav')

    @override_settings(MEDIA_ROOT=MEDIA_ROOT)
    def test_new_uploads_go_to_day_folder(self):
        audio = make_call(CALL_A)
        # The shared test media folder may already hold this name: Django adds a suffix then.
        self.assertRegex(audio.file.name, r'^audio/2026/09/25/out-375290000103-308-20260925-093215-1790000001(_\w+)?\.100\.wav$')
        self.assertTrue(Path(audio.file.path).exists())

    def test_relocate_old_flat_files(self):
        with tempfile.TemporaryDirectory() as root:
            (Path(root) / 'audio').mkdir()
            (Path(root) / 'audio' / CALL_A).write_bytes(b'A')
            (Path(root) / 'audio' / 'manual.wav').write_bytes(b'M')
            # a file of the same name already in the day folder
            (Path(root) / 'audio/2026/09/25').mkdir(parents=True)
            (Path(root) / 'audio/2026/09/25' / CALL_B).write_bytes(b'old')
            (Path(root) / 'audio' / CALL_B).write_bytes(b'B')
            a = AudioFile.objects.create(file=f'audio/{CALL_A}')
            b = AudioFile.objects.create(file=f'audio/{CALL_B}')
            manual = AudioFile.objects.create(file='audio/manual.wav',
                                              call_started_at=timezone.make_aware(datetime(2026, 9, 27, 10, 0)))
            missing = AudioFile.objects.create(file='audio/out-375290000999-301-20260925-100000-1790000099.199.wav')

            self.assertEqual(audio_paths.relocate(AudioFile, root), 3)
            a.refresh_from_db(); b.refresh_from_db(); manual.refresh_from_db(); missing.refresh_from_db()
            self.assertEqual(a.file.name, f'audio/2026/09/25/{CALL_A}')
            self.assertEqual((Path(root) / a.file.name).read_bytes(), b'A')
            self.assertEqual(b.file.name, f'audio/2026/09/25/{Path(CALL_B).stem}_1.wav')  # name was taken
            self.assertEqual(manual.file.name, 'audio/2026/09/27/manual.wav')
            self.assertEqual(missing.file.name, 'audio/out-375290000999-301-20260925-100000-1790000099.199.wav')
            self.assertFalse((Path(root) / 'audio' / CALL_A).exists())
            self.assertEqual(audio_paths.relocate(AudioFile, root), 0)  # running again changes nothing

    @override_settings(MEDIA_ROOT=MEDIA_ROOT)
    def test_pbx_download_goes_to_day_folder(self):
        AppSettings.objects.create(pk=1, min_duration=30)
        run_sync(PbxSync.objects.create(day=date(2026, 9, 25)), client=FakePbx([rec(CALL_B)]))
        self.assertRegex(AudioFile.objects.get().file.name,
                         r'^audio/2026/09/25/out-375290000104-308-20260925-095111-1790000002(_\w+)?\.101\.wav$')


@override_settings(MEDIA_ROOT=MEDIA_ROOT)
class SortAndStatsTests(TestCase):
    def setUp(self):
        Interviewer.objects.create(extension='308', name='Марина')
        Interviewer.objects.create(extension='306', name='Елена')
        self.a = make_call(CALL_A, duration=120)   # 308, 09:32:15
        self.b = make_call(CALL_B, duration=60)    # 308, 09:51:11
        self.c = make_call(CALL_C, duration=300)   # 306, 09:18:17
        CallReview.objects.create(audio=self.a, interviewer='Марина 308', result='ок', completed=True)
        CallReview.objects.create(audio=self.c, interviewer='Елена 306', result='ошибка', completed=True)
        CallReview.objects.create(audio=self.b, result='брак', completed=False)  # draft, no interviewer

    def phones(self, sort):
        r = self.client.get(reverse('player:index'), {'day': '2026-09-25', 'sort': sort})
        return [a.phone for a in r.context['calls']], r

    def test_sort_by_time_and_interviewer(self):
        self.assertEqual(self.phones('time')[0], ['375290000105', '375290000103', '375290000104'])
        self.assertEqual(self.phones('-time')[0], ['375290000104', '375290000103', '375290000105'])
        self.assertEqual(self.phones('interviewer')[0], ['375290000105', '375290000103', '375290000104'])
        self.assertEqual(self.phones('-duration')[0], ['375290000105', '375290000103', '375290000104'])
        self.assertEqual(self.phones('nonsense')[0], self.phones('time')[0])

    def test_interviewer_column_uses_names(self):
        _, r = self.phones('time')
        self.assertEqual([a.interviewer_label for a in r.context['calls']], ['Елена 306', 'Марина 308', 'Марина 308'])

    def test_header_links_toggle_direction(self):
        _, r = self.phones('interviewer')
        cols = {c['label']: c for c in r.context['columns']}
        self.assertEqual(cols['Интервьюер']['arrow'], '↑')
        self.assertTrue(cols['Интервьюер']['url'].endswith('day=2026-09-25&sort=-interviewer'))
        self.assertTrue(cols['Время']['url'].endswith('sort=time'))
        self.assertNotIn('url', cols['Телефон'])

    def test_day_stats(self):
        _, r = self.phones('time')
        rows = {row.interviewer: row for row in r.context['day_stats']}
        marina, elena = rows['Марина 308'], rows['Елена 306']
        self.assertEqual((marina.calls, marina.reviewed, marina.ok, marina.rejects), (2, 1, 1, 0))  # draft not counted
        self.assertEqual(marina.seconds, 180)
        self.assertEqual((elena.errors, elena.error_share), (1, 100))
        total = r.context['day_total']
        self.assertEqual((total.calls, total.reviewed, round(total.error_share)), (3, 2, 50))
        self.assertContains(r, 'Интервьюеры за 25.09.2026')

    def test_stats_page_period(self):
        make_call('out-375290000301-306-20260801-100000-1780000001.300.wav', duration=90)  # outside
        r = self.client.get(reverse('player:stats'), {'from': '2026-09-01', 'to': '2026-09-30'})
        self.assertEqual(r.context['total'].calls, 3)
        self.assertContains(r, 'Марина 308')
        r = self.client.get(reverse('player:stats'), {'from': '2026-09-30', 'to': '2026-08-01'})  # swapped
        self.assertEqual(r.context['total'].calls, 4)

    def test_error_share_ignores_reviews_without_result(self):
        CallReview.objects.filter(audio=self.c).update(result='')  # reviewed, no result chosen
        _, r = self.phones('time')
        elena = next(row for row in r.context['day_stats'] if row.interviewer == 'Елена 306')
        self.assertEqual((elena.reviewed, elena.error_share), (1, None))

    def test_hms_filter(self):
        from player.templatetags.player_extras import hms
        self.assertEqual([hms(5), hms(312), hms(3723.4), hms(None)], ['0:05', '5:12', '1:02:03', ''])


SURVEY_PAGE = """<html><body>
<input class="w-50 max-h-row" id="RespName" type="text" value="Мария">
<select name="RespCity" id="RespCity"><option value="1">Брест</option><option value="4" selected>Гомель</option>
<option value="7">Минск</option></select>
<input class="w-50 max-h-row" id="RespAge" type="Number" value="56">
<select id="RespGender"><option value="1" selected>Женщина</option><option value="2">Мужчина</option></select>
<select id="RespEdStatus"><option class="opt" value="3">Среднее специальное (техникум, колледж и т. д.)</option>
<option class="opt" value="4" selected>Высшее (в том числе магистратура, аспирантура, соискательство)</option></select>
<select id="RespWorkStatus"><option value="3" selected>Служащий</option></select>
<select id="RespIncome"><option value="4">75 % и более</option><option value="3" selected>50-75%</option></select>
<select id="RespRecall"><option value="1" selected>Да</option></select>
<a href="/admin/Reports/delete200935">Удалить анкету</a>
<script src="/assets/admin/Reports/js/UpdateScript.js" data-Members="258738"></script>
</body></html>"""
GET_MEMBER = {  # the shape of /admin/Reports/getMember, made-up person
    'MemberName': 'Мария', 'DateOfBirth': '56', 'GenderID': 1, 'CityID': 4, 'MembersSurveyID': '258738',
    'MemberID': '99999901', 'SurveyDate': {'date': '2026-09-24 00:00:00.000000', 'timezone_type': 3},
    'SurveyDataJSON': {'WeekListening': {'flag': 'No'}, 'MemberID': '99999901'},
    'OPUserID': 7, 'OPTime': '25-09-2026 09:34:01',
}


class CrmSurveyReadingTests(TestCase):
    def test_parse_survey_page(self):
        profile, members_id = crm.parse_survey_page(SURVEY_PAGE)
        self.assertEqual(members_id, '258738')
        self.assertEqual(profile['Город'], 'Гомель')
        self.assertEqual(profile['Возраст'], '56')
        self.assertEqual(profile['Образование'], 'Высшее (в том числе магистратура, аспирантура, соискательство)')
        self.assertEqual(profile['Доля дохода на продукты питания'], '50-75%')
        self.assertNotIn('Мария', json.dumps(profile, ensure_ascii=False))  # the name is not read

    def test_fetch_answers_without_personal_data(self):
        client = mock.Mock()
        client.url.side_effect = lambda path: 'http://crm.local/' + path
        client.request.side_effect = [crm_reply(SURVEY_PAGE), crm_reply('{}', json_data=GET_MEMBER)]
        survey = crm.parse_survey(CRM_FIND[1])
        answers = crm.fetch_answers(client, survey)
        post = client.request.call_args_list[1]
        self.assertEqual(post.args, ('POST', 'admin/Reports/getMember'))
        self.assertEqual(post.kwargs['data'], {'MembersSurveyID': '258738'})
        self.assertEqual(answers['День «вчера»'], '24.09.2026')
        self.assertEqual(answers['Ответы о слушании'], {'WeekListening': {'flag': 'No'}})
        text = json.dumps(answers, ensure_ascii=False)
        self.assertNotIn('Мария', text)
        self.assertNotIn('99999901', text)


@override_settings(MEDIA_ROOT=MEDIA_ROOT, GEMINI_API_KEY='test-key', GEMINI_MODEL='gemini-test')
class SurveyCheckRunTests(TestCase):
    RESULT = compare_module.Comparison(
        discrepancies=[
            compare_module.Discrepancy(field='Возраст', survey_value='26', call_value='26 неполных', time='0:31',
                                       severity='ошибка', comment='следовало указать 25'),
            compare_module.Discrepancy(field='Город', survey_value='Гомель', call_value='Гомель?', time='',
                                       severity='странное', comment='уточнить'),
        ],
        summary='Есть ошибка в возрасте.',
    )

    def setUp(self):
        AppSettings.objects.update_or_create(pk=1, defaults={'crm_url': 'http://crm.local', 'crm_username': 'sv'})
        self.audio = make_call(CALL_A, status=AudioFile.Status.DONE)
        Segment.objects.create(audio=self.audio, index=0, start=31, end=33, text='Мне 26 неполных.')
        self.crm_client = mock.Mock()
        self.crm_client.find.return_value = [crm.parse_survey(row) for row in CRM_FIND]

    def run_check(self, result=None):
        compare_module.queue(self.audio)
        item = self.audio.comparison
        fake = FakeGemini(result or self.RESULT)
        with mock.patch('player.crm.fetch_answers', return_value={'Профиль': {'Возраст': '26'}}):
            compare_module.run_comparison(item, gemini_client=fake, crm_client=self.crm_client)
        item.refresh_from_db()
        return item, fake

    def test_discrepancies_saved(self):
        item, fake = self.run_check()
        self.assertEqual((item.status, item.survey_id), ('done', 200))
        self.assertEqual(item.discrepancies[0]['comment'], 'следовало указать 25')
        self.assertEqual(item.discrepancies[1]['severity'], 'проверить')  # unknown severity normalised
        sent = fake.calls[0]['contents']
        self.assertIn('[00:31] Мне 26 неполных.', sent)
        self.assertNotIn(self.audio.phone, sent + fake.calls[0]['config'].system_instruction)

    def test_no_survey_for_the_call(self):
        self.crm_client.find.return_value = []
        item, fake = self.run_check()
        self.assertEqual(item.status, 'error')
        self.assertIn('не найдена', item.error)
        self.assertEqual(fake.calls, [])

    def test_crm_down_is_retried(self):
        self.crm_client.find.side_effect = crm.CrmError('CRM недоступна: timeout')
        compare_module.queue(self.audio)
        with self.assertRaises(analysis.RetryLater):
            compare_module.run_comparison(self.audio.comparison, gemini_client=FakeGemini(self.RESULT),
                                          crm_client=self.crm_client)

    def test_endpoint_and_panel(self):
        r = self.client.get(self.audio.get_absolute_url())
        self.assertContains(r, 'id="compare-panel"')
        r = self.client.post(reverse('player:comparison', args=[self.audio.pk]))
        self.assertEqual(r.json()['comparison']['status'], 'pending')

    def test_worker_queues_check_after_transcription(self):
        pending = make_call(CALL_B)

        def fake_transcribe(audio):
            Segment.objects.create(audio=audio, index=0, start=0, end=2, text='Минск.')
            AudioFile.objects.filter(pk=audio.pk).update(status='done')

        with mock.patch('player.worker.transcribe', side_effect=fake_transcribe), \
                mock.patch('player.worker.run_analysis_step', return_value=False):
            worker.run_once()
        self.assertEqual(pending.comparison.status, CallComparison.Status.PENDING)
