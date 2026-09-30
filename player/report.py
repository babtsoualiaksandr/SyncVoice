"""Daily Excel report in the controllers' format («Отчет записи ДД.ММ.ГГГГ.xlsx»)."""
import io
from datetime import date

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from .models import AudioFile

# (header, width) for columns A..P, exactly as in the controllers' template.
COLUMNS = [
    (None, 4.2),
    ('Дата контроля', 10.7),
    ('№ телефона', 18.0),
    ('Дата опроса  Время', 21.0),
    ('Интервьюер: код (например, Ольга Инна Инна Инна 303 Марина) и Имя', 19.5),
    ('ID респондента', 5.8),
    ('Город проживания', 12.2),
    ('Город слушания', 10.5),
    ('Радиостанции / не слушал', 17.0),
    ('Ошибки ', 8.5),
    ('Комментарий к ошибкам', 29.3),
    ('Результат: брак/ошибки/ок', 16.2),
    ('Прим', 12.5),
    ('Продолжительность сек', 9.2),
    ('Контроль', 15.0),
    (None, 9.2),
]
LAST_FILTER_COLUMN = 'O'


def report_filename(day: date) -> str:
    return f'Отчет записи {day:%d.%m.%Y}.xlsx'


def calls_for_day(day: date):
    return (
        AudioFile.objects.filter(call_started_at__date=day)
        .select_related('review')
        .order_by('call_started_at')
    )


def build_report(day: date) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Лист1'

    for i, (header, width) in enumerate(COLUMNS, 1):
        letter = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[letter].width = width
        if header:
            cell = ws.cell(row=1, column=i, value=header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical='top')

    rows = []
    for audio in calls_for_day(day):
        review = getattr(audio, 'review', None)
        interviewer = review.interviewer if review and review.interviewer else audio.default_interviewer()
        rows.append((interviewer, audio, review))
    # Grouped by interviewer, then by call time — as controllers fill it in.
    rows.sort(key=lambda r: (r[0], r[1].call_started_at))

    for row_no, (interviewer, audio, review) in enumerate(rows, 2):
        started = timezone.localtime(audio.call_started_at).replace(tzinfo=None)
        values = {
            'B': review.review_date if review else None,
            'C': audio.phone,
            'D': started,
            'E': interviewer,
            'N': round(audio.duration) if audio.duration else None,
        }
        if review:
            values.update({
                'F': review.respondent_id,
                'G': review.city,
                'H': review.listen_city,
                'I': review.stations,
                'J': 'есть' if review.has_errors else '',
                'K': review.error_comment,
                'L': review.result,
                'M': review.note,
                'O': review.controller,
                'P': review.fixed,
            })
        for col, value in values.items():
            if value in (None, ''):
                continue
            ws[f'{col}{row_no}'] = value
        ws[f'B{row_no}'].number_format = 'dd.mm.yyyy'
        ws[f'C{row_no}'].number_format = '@'
        ws[f'D{row_no}'].number_format = 'yyyy-mm-dd h:mm:ss'

    ws.freeze_panes = 'D2'
    ws.auto_filter.ref = f'A1:{LAST_FILTER_COLUMN}{max(len(rows) + 1, 2)}'

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
